from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.datavalidation import DataValidation

from app_reports.excel import LOGO

from .models import Sender

FONT = "Arial"
line = Side(style="thin", color="000000")
box = Border(left=line, right=line, top=line, bottom=line)

# (подпись, высота строки в пунктах): описание и требования пишут развёрнуто
FIELDS = [
    ("Изделие", 24),
    ("Номер изделия и цвет", 24),
    ("Деталь", 24),
    ("Описание претензии", 90),
    ("Данные с упаковки", 36),
    ("№ ТН и дата отгрузки", 24),
    ("Наименование отправителя\n(выберите из списка)", 30),
    ("Цена по накладной", 24),
    ("Дата отгрузки конечному покупателю", 24),
    ("Требования покупателя", 72),
]


def build_blank() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт-претензия"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 58

    logo = XlImage(str(LOGO))
    logo.width = logo.height = 56
    ws.add_image(logo, "A1")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    ws["B1"] = "Мебельная фабрика «Миларум»"
    ws["B1"].font = Font(name=FONT, size=11, bold=True)
    ws["B2"] = "Отдел продаж: 8 (800) 511-15-10"
    ws["B2"].font = Font(name=FONT, size=9)
    for cell in ("B1", "B2"):
        ws[cell].alignment = Alignment(horizontal="right", vertical="center")

    ws["A4"] = "Дата:"
    ws["A4"].font = Font(name=FONT, size=10)
    ws["A4"].alignment = Alignment(horizontal="right", vertical="center")
    ws["B4"].border = Border(bottom=line)
    ws["B4"].number_format = "DD.MM.YYYY"
    ws["B4"].font = Font(name=FONT, size=10)

    for row, text, font in (
        (6, "АКТ-ПРЕТЕНЗИЯ", Font(name=FONT, size=15, bold=True)),
        (7, "Данная форма заполняется клиентом", Font(name=FONT, size=10)),
        (8, "ОБЯЗАТЕЛЬНО К АКТУ-ПРЕТЕНЗИИ ПРИЛАГАТЬ ФОТО БРАКА!", Font(name=FONT, size=10, bold=True, color="C00000")),
    ):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1, text).font = font
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[6].height = 26

    first = 10
    for i, (label, height) in enumerate(FIELDS):
        row = first + i
        key = ws.cell(row, 1, label)
        key.font = Font(name=FONT, size=10)
        value = ws.cell(row, 2)
        value.font = Font(name=FONT, size=10)
        for cell in (key, value):
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        ws.row_dimensions[row].height = height

    senders = DataValidation(
        type="list",
        formula1='"' + ",".join(Sender.labels) + '"',
        allow_blank=True,
        showErrorMessage=False,
        promptTitle="Отправитель",
        prompt="Выберите из списка",
    )
    ws.add_data_validation(senders)
    senders.add(ws.cell(first + 6, 2))
    ws.cell(first + 8, 2).number_format = "DD.MM.YYYY"

    sign = first + len(FIELDS) + 2
    ws.row_dimensions[sign].height = 24
    ws.cell(sign, 1, "Подпись клиента / расшифровка:").font = Font(name=FONT, size=10)
    ws.cell(sign, 1).alignment = Alignment(vertical="bottom")
    ws.cell(sign, 2).border = Border(bottom=line)

    ws.print_area = f"A1:B{sign}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
