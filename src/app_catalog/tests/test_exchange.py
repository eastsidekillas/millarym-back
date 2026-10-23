from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from app_catalog.models import Category, Fabric, KitOption, ProductModel, Team
from app_production.tests.base import Fixtures


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


class ExchangeTest(Fixtures, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.make_fixtures()

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.director)

    def export(self):
        r = self.client.get("/api/v1/catalog/export/")
        self.assertEqual(r.status_code, 200)
        return load_workbook(BytesIO(r.content))

    def upload(self, wb, dry_run=False, user=None):
        out = BytesIO()
        wb.save(out)
        client = self.client
        if user:
            client = APIClient()
            client.force_authenticate(user)
        url = "/api/v1/catalog/import/" + ("?dry_run=1" if dry_run else "")
        return client.post(url, {"file": SimpleUploadedFile("c.xlsx", out.getvalue())}, format="multipart")

    def sheet(self, data, title):
        return next(s for s in data["sheets"] if s["title"] == title)

    def test_roundtrip_without_changes_creates_nothing(self):
        before = (ProductModel.objects.count(), Category.objects.count(), Fabric.objects.count(), Team.objects.count())
        data = self.upload(self.export()).json()
        self.assertTrue(data["applied"])
        self.assertTrue(all(not s["created"] and not s["updated"] and not s["errors"] for s in data["sheets"]))
        after = (ProductModel.objects.count(), Category.objects.count(), Fabric.objects.count(), Team.objects.count())
        self.assertEqual(before, after)

    def test_rename_by_id_new_row_and_kit(self):
        wb = self.export()
        ws = wb["Модели"]
        rows = sheet_rows(ws)
        boston = next(i for i, r in enumerate(rows) if r[1] == "Бостон") + 1
        ws.cell(boston, 2, "Бостон Люкс")
        ws.cell(boston, 4, "Подушки 3 шт.; Подлокотники мягкие")
        ws.append([None, "Новая модель", "Прямые диваны", "", "да"])
        florence = next(i for i, r in enumerate(rows) if r[1] == "Флоренция") + 1
        ws.cell(florence, 1, None)  # ID стёрли — найдётся по названию
        ws.cell(florence, 2, "флоренция")
        ws.cell(florence, 5, "нет")
        wb["Бригады"].append([None, "Бригада Петрова", "Петров", 5, "да"])

        preview = self.upload(wb, dry_run=True).json()
        self.assertFalse(preview["applied"])
        self.assertEqual(self.sheet(preview, "Модели")["created"], ["Новая модель"])
        self.assertFalse(ProductModel.objects.filter(name="Новая модель").exists())

        data = self.upload(wb).json()
        self.assertTrue(data["applied"], data)
        self.model.refresh_from_db()
        self.assertEqual((self.model.name, self.model.kit_options), ("Бостон Люкс", ["Подушки 3 шт.", "Подлокотники мягкие"]))
        self.other_model.refresh_from_db()
        self.assertFalse(self.other_model.active)
        self.assertEqual(ProductModel.objects.filter(name__iexact="флоренция").count(), 1)
        self.assertTrue(Team.objects.filter(name="Бригада Петрова", foreman="Петров").exists())

    def test_errors_change_nothing(self):
        wb = self.export()
        ws = wb["Модели"]
        ws.append([None, "Модель А", "Нет такой категории", "", "да"])
        ws.append([99999, "Модель Б", "Прямые диваны", "", "да"])
        ws.append([None, "Модель В", "Прямые диваны", "подушки, валики", "да"])
        wb["Категории"].append([None, "Новая категория", 1, "да"])
        data = self.upload(wb).json()
        self.assertFalse(data["applied"])
        errors = " ".join(self.sheet(data, "Модели")["errors"])
        self.assertIn("нет категории", errors)
        self.assertIn("нет записи с ID 99999", errors)
        self.assertIn("без запятых", errors)
        self.assertFalse(Category.objects.filter(name="Новая категория").exists())

    def test_file_from_other_base_matches_by_name(self):
        wb = self.export()
        ws = wb["Модели"]
        for row in ws.iter_rows(min_row=2):
            row[0].value = (row[0].value or 0) + 1000  # чужие ID
        ws.append([5000, "Модель с другого сервера", "Прямые диваны", "", "да"])
        wb["служебное"]["B1"] = "чужая"
        data = self.upload(wb).json()
        self.assertTrue(data["applied"], data)
        self.assertTrue(data["foreign"])
        self.assertEqual(self.sheet(data, "Модели")["created"], ["Модель с другого сервера"])
        self.assertEqual(ProductModel.objects.filter(name="Бостон").count(), 1)

    def test_rights(self):
        wb = self.export()
        wb["Ткани"].append([None, "Велюр 99", "да"])
        wb["Категории"].append([None, "Пуфики", 1, "да"])
        data = self.upload(wb, user=self.cutter).json()  # ткани — можно, справочники — нет
        self.assertTrue(data["applied"])
        self.assertTrue(Fabric.objects.filter(name="Велюр 99").exists())
        self.assertFalse(Category.objects.filter(name="Пуфики").exists())
        self.assertIn("нет прав", self.sheet(data, "Категории")["skipped"])
        self.assertEqual(self.upload(wb, user=self.manager).status_code, 403)

    def test_new_category_used_by_new_model(self):
        wb = self.export()
        wb["Категории"].append([None, "Кресла-качалки", 9, "да"])
        wb["Модели"].append([None, "Качалка", "Кресла-качалки", "", ""])
        wb["Комплектация"].append([None, "Чехол на подголовник", None, None])
        self.assertTrue(self.upload(wb).json()["applied"])
        self.assertEqual(ProductModel.objects.get(name="Качалка").category.name, "Кресла-качалки")
        self.assertTrue(KitOption.objects.get(name="Чехол на подголовник").active)
