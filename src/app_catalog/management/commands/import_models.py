import csv
import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from app_catalog.models import Category, ProductModel

NAME_HEADERS = {"название", "наименование", "модель", "изделие", "name"}
CATEGORY_HEADERS = {"категория", "category", "группа", "тип"}


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def read_rows(path: Path, sheet: str | None) -> list[list[str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        ws = book[sheet] if sheet else book.worksheets[0]
        return [[clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise CommandError("Не удалось прочитать файл: сохраните его в UTF-8 или Windows-1251.")
    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
    return [[clean(c) for c in row] for row in csv.reader(text.splitlines(), dialect)]


class Command(BaseCommand):
    help = "Загрузить справочник моделей из Excel или CSV (колонки «Название» и «Категория»)."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", help="лист Excel (по умолчанию первый)")
        parser.add_argument("--category", help="категория для всех строк, если колонки категории нет")
        parser.add_argument("--create-categories", action="store_true", help="заводить неизвестные категории")
        parser.add_argument("--update-category", action="store_true", help="менять категорию у уже заведённых моделей")
        parser.add_argument("--dry-run", action="store_true", help="только показать, что будет сделано")

    def handle(self, *args, path, sheet, category, create_categories, update_category, dry_run, **options):
        file = Path(path)
        if not file.exists():
            raise CommandError(f"Нет файла {file}.")
        rows = [r for r in read_rows(file, sheet) if any(r)]
        if not rows:
            raise CommandError("Файл пустой.")

        header = [h.lower() for h in rows[0]]
        name_col = next((i for i, h in enumerate(header) if h in NAME_HEADERS), None)
        cat_col = next((i for i, h in enumerate(header) if h in CATEGORY_HEADERS), None)
        if name_col is None:
            raise CommandError(f"Не нашли колонку с названием. Заголовки: {', '.join(rows[0])}.")
        if cat_col is None and not category:
            raise CommandError("Нет колонки «Категория» — укажите категорию для всех: --category «Прямые диваны».")

        categories = {c.name.lower(): c for c in Category.objects.all()}
        existing = {m.name.lower(): m for m in ProductModel.objects.select_related("category")}
        report = {"created": [], "exists": [], "recategorized": [], "duplicates": [], "no_category": [], "new_categories": []}
        seen = set()

        with transaction.atomic():
            for line, row in enumerate(rows[1:], start=2):
                name = row[name_col] if name_col < len(row) else ""
                cat_name = (row[cat_col] if cat_col is not None and cat_col < len(row) else "") or (category or "")
                if not name:
                    continue
                key = name.lower()
                if key in seen:
                    report["duplicates"].append(f"{line}: {name}")
                    continue
                seen.add(key)

                cat = categories.get(cat_name.lower())
                if cat is None and cat_name and create_categories:
                    cat = Category.objects.create(name=cat_name, sort=(Category.objects.count() + 1) * 10)
                    categories[cat_name.lower()] = cat
                    report["new_categories"].append(cat_name)
                if cat is None:
                    report["no_category"].append(f"{line}: {name} ({cat_name or 'категория не указана'})")
                    continue

                model = existing.get(key)
                if model:
                    if update_category and model.category_id != cat.pk:
                        model.category = cat
                        model.save(update_fields=["category"])
                        report["recategorized"].append(f"{name}: {cat.name}")
                    else:
                        report["exists"].append(name)
                    continue
                existing[key] = ProductModel.objects.create(name=name, category=cat)
                report["created"].append(name)

            if dry_run:
                transaction.set_rollback(True)

        self.print_report(report, dry_run)

    def print_report(self, report, dry_run):
        out = self.stdout
        prefix = "Проверка без записи. " if dry_run else ""
        out.write(self.style.SUCCESS(
            f"{prefix}Новых моделей: {len(report['created'])}, уже были: {len(report['exists'])}, "
            f"сменили категорию: {len(report['recategorized'])}."
        ))
        if report["new_categories"]:
            out.write(f"Новые категории: {', '.join(report['new_categories'])}")
        for title, key in (
            ("Пропущены — неизвестная категория (добавьте её или запустите с --create-categories)", "no_category"),
            ("Пропущены — повтор в файле", "duplicates"),
        ):
            if report[key]:
                out.write(self.style.WARNING(f"{title}: {len(report[key])}"))
                for row in report[key]:
                    out.write(f"  {row}")
