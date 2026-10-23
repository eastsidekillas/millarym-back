"""Справочники в Excel: выгрузить, поправить, загрузить обратно.

Строка находится по колонке ID — так переименование меняет запись, а не заводит новую. Без ID ищется по названию
(без регистра), не нашлась — создаётся. Строки, которых нет в файле, не трогаются: удалять — только в программе.
Ошибка хоть в одной строке — не меняется ничего.
"""

from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Category, Fabric, KitOption, ProductModel, Team

YES = {"да", "yes", "1", "true", "+", "активна", "активен"}
NO = {"нет", "no", "0", "false", "-", "скрыта", "скрыт"}
HEAD_FILL = PatternFill("solid", fgColor="E8EFE6")
KIT_SEPARATOR = ";"


@dataclass(frozen=True)
class Col:
    title: str
    attr: str
    kind: str = "text"  # text | int | bool | category | kit
    width: int = 30


@dataclass(frozen=True)
class Sheet:
    title: str
    model: type
    cols: tuple[Col, ...]
    right: str = "catalog"


SHEETS = [
    Sheet("Категории", Category, (Col("Название", "name"), Col("Порядок", "sort", "int", 10), Col("Активна", "active", "bool", 10))),
    Sheet(
        "Модели",
        ProductModel,
        (
            Col("Название", "name", width=36),
            Col("Категория", "category", "category", 24),
            Col("Своя комплектация (через ;)", "kit_options", "kit", 44),
            Col("Активна", "active", "bool", 10),
        ),
    ),
    Sheet("Комплектация", KitOption, (Col("Название", "name"), Col("Порядок", "sort", "int", 10), Col("Активна", "active", "bool", 10))),
    Sheet(
        "Бригады",
        Team,
        (Col("Бригада", "name"), Col("Бригадир", "foreman"), Col("Порядок", "sort", "int", 10), Col("Активна", "active", "bool", 10)),
    ),
    Sheet("Ткани", Fabric, (Col("Название", "name", width=36), Col("Активна", "active", "bool", 10)), right="fabrics"),
]

HELP = [
    "Как заполнять",
    "",
    "Колонку ID не меняйте: по ней программа узнаёт запись. Поменяли название — запись переименуется, новой не будет.",
    "Новая запись — новая строка с пустым ID.",
    "Строку удалили из файла — в программе запись останется. Скрыть — «нет» в колонке «Активна», удалить — в программе.",
    "Категорию модели пишите точно как на листе «Категории» (можно выбрать из списка).",
    "Своя комплектация модели — через точку с запятой: «Подушки-подголовники 3 шт.; Подлокотники мягкие». Без запятых.",
    "Фото тканей и моделей в файл не попадают и при загрузке не меняются.",
    "Перед загрузкой программа покажет, что изменится. Если есть ошибки — не изменится ничего.",
]


def _cell_value(obj, col: Col):
    value = getattr(obj, col.attr)
    if col.kind == "bool":
        return "да" if value else "нет"
    if col.kind == "category":
        return value.name
    if col.kind == "kit":
        return f"{KIT_SEPARATOR} ".join(value or [])
    return value


def build_xlsx() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.title)
        ws.append(["ID", *(c.title for c in sheet.cols)])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
        ws.column_dimensions["A"].width = 8
        for i, col in enumerate(sheet.cols, start=2):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = col.width
        qs = sheet.model.objects.all()
        if sheet.model is ProductModel:
            qs = qs.select_related("category")
        for obj in qs:
            ws.append([obj.pk, *(_cell_value(obj, c) for c in sheet.cols)])
        ws.freeze_panes = "B2"
        for i, col in enumerate(sheet.cols, start=2):
            letter = ws.cell(1, i).column_letter
            if col.kind == "bool":
                dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
            elif col.kind == "category":
                dv = DataValidation(type="list", formula1="=Категории!$B$2:$B$500", allow_blank=False)
            else:
                continue
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}2000")
    ws = wb.create_sheet("Как заполнять")
    for line in HELP:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


@dataclass
class SheetReport:
    title: str
    created: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    unchanged: int = 0
    errors: list[str] = field(default_factory=list)
    skipped: str = ""


class ImportError_(Exception):
    pass


def _norm(text) -> str:
    return " ".join(str(text or "").split())


def _key(text: str) -> str:
    return _norm(text).lower().replace("ё", "е")


def _parse(col: Col, raw, categories: dict[str, Category]):
    if col.kind == "text":
        return _norm(raw)
    if col.kind == "int":
        if raw in (None, ""):
            return None
        try:
            value = int(float(str(raw).replace(",", ".")))
        except ValueError:
            raise ImportError_(f"«{col.title}»: нужно число, а не «{raw}»")
        if not 0 <= value <= 32767:
            raise ImportError_(f"«{col.title}»: число от 0 до 32767")
        return value
    if col.kind == "bool":
        text = _key(raw)
        if not text:
            return None
        if text in YES:
            return True
        if text in NO:
            return False
        raise ImportError_(f"«{col.title}»: «да» или «нет», а не «{raw}»")
    if col.kind == "category":
        name = _norm(raw)
        if not name:
            raise ImportError_("не указана категория")
        category = categories.get(_key(name))
        if category is None:
            raise ImportError_(f"нет категории «{name}» — добавьте её на лист «Категории»")
        return category
    if col.kind == "kit":
        options = []
        for part in str(raw or "").replace("\n", KIT_SEPARATOR).split(KIT_SEPARATOR):
            part = _norm(part)
            if not part:
                continue
            if "," in part:
                raise ImportError_(f"комплектация без запятых: «{part}»")
            if _key(part) not in (_key(o) for o in options):
                options.append(part)
        return options
    raise AssertionError(col.kind)


def _show(col: Col, value) -> str:
    if col.kind == "bool":
        return "да" if value else "нет"
    if col.kind == "category":
        return value.name if value else "—"
    if col.kind == "kit":
        return f"{KIT_SEPARATOR} ".join(value) or "—"
    return str(value) if value not in (None, "") else "—"


def _import_sheet(ws, sheet: Sheet, report: SheetReport, categories: dict[str, Category]) -> None:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    header = [_norm(h) for h in rows[0]]
    expected = ["ID", *(c.title for c in sheet.cols)]
    missing = [t for t in expected if t not in header]
    if missing:
        report.errors.append(f"нет колонок: {', '.join(missing)} — выгрузите файл заново и правьте его")
        return
    index = {t: header.index(t) for t in expected}
    name_col = sheet.cols[0]
    existing = {obj.pk: obj for obj in sheet.model.objects.all()}
    by_name = {_key(obj.name): obj for obj in existing.values()}
    seen_names: dict[str, int] = {}
    seen_ids: set[int] = set()

    for number, row in enumerate(rows[1:], start=2):
        def value(title):
            i = index[title]
            return row[i] if i < len(row) else None

        if all(v in (None, "") for v in row):
            continue
        try:
            raw_id = value("ID")
            pk = None
            if raw_id not in (None, ""):
                try:
                    pk = int(float(str(raw_id)))
                except ValueError:
                    raise ImportError_(f"ID «{raw_id}» — должно быть число или пусто для новой записи")
            data = {c.attr: _parse(c, value(c.title), categories) for c in sheet.cols}
            name = data[name_col.attr]
            if not name:
                raise ImportError_(f"пустое «{name_col.title}»")
            key = _key(name)
            if key in seen_names:
                raise ImportError_(f"«{name}» уже есть в строке {seen_names[key]}")
            seen_names[key] = number

            if pk is not None:
                obj = existing.get(pk)
                if obj is None:
                    raise ImportError_(f"нет записи с ID {pk} — для новой оставьте ID пустым")
                if pk in seen_ids:
                    raise ImportError_(f"ID {pk} повторяется")
                seen_ids.add(pk)
            else:
                obj = by_name.get(key)
                if obj is not None and obj.pk in seen_ids:
                    raise ImportError_(f"«{name}» уже есть в файле под ID {obj.pk}")
            if obj is not None:
                seen_ids.add(obj.pk)
                other = by_name.get(key)
                if other is not None and other.pk != obj.pk:
                    raise ImportError_(f"название «{name}» уже у другой записи (ID {other.pk})")
                changes = []
                for col in sheet.cols:
                    new = data[col.attr]
                    if new is None:  # пустые «Активна» и «Порядок» — оставить как было
                        continue
                    old = getattr(obj, col.attr)
                    if (old.pk if col.kind == "category" else old) != (new.pk if col.kind == "category" else new):
                        changes.append(f"{col.title}: {_show(col, old)} → {_show(col, new)}")
                        setattr(obj, col.attr, new)
                if changes:
                    obj.save()
                    report.updated.append(f"{name}: " + "; ".join(changes))
                else:
                    report.unchanged += 1
            else:
                obj = sheet.model(**{k: v for k, v in data.items() if v is not None})
                obj.save()
                existing[obj.pk] = obj
                by_name[key] = obj
                seen_ids.add(obj.pk)
                report.created.append(name)
            if sheet.model is Category:
                categories[_key(obj.name)] = obj
        except ImportError_ as exc:
            report.errors.append(f"строка {number}: {exc}")


def import_xlsx(content: bytes, can, dry_run: bool) -> list[SheetReport]:
    """can(right) — есть ли у пользователя право. Ошибка где угодно или проверка — откат всего."""
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Не удалось открыть файл — нужен .xlsx, выгруженный из «Справочников».")
    reports = []
    known = [s for s in SHEETS if s.title in wb.sheetnames]
    if not known:
        raise ValueError("В файле нет листов справочников — выгрузите файл из программы и правьте его.")
    with transaction.atomic():
        categories = {_key(c.name): c for c in Category.objects.all()}
        for sheet in known:
            report = SheetReport(sheet.title)
            reports.append(report)
            if not can(sheet.right):
                report.skipped = "нет прав менять этот справочник — лист пропущен"
                continue
            _import_sheet(wb[sheet.title], sheet, report, categories)
        if dry_run or any(r.errors for r in reports):
            transaction.set_rollback(True)
    return reports
