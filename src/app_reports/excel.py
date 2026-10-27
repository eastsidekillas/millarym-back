import re
from io import BytesIO
from pathlib import Path

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image, ImageOps

from .reports import Result

LOGO = Path(__file__).parent / "assets" / "logo.png"
FONT = "Arial"
INK = "1E2223"
MUTED = "545859"
ACCENT = "006342"
LINE = "DCE3DA"
HEAD_FILL = "EEF3EC"
MARK_FILL = {"warn": "FFF4E0", "crit": "FBE9E7"}
PHOTO = 44
PHOTO_GAP = 4
GROUP_PHOTO = (120, 60)
HEADER_ROW = 6

thin = Side(style="thin", color=LINE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _photo(path: str, size: tuple[int, int] = (PHOTO, PHOTO), whole: bool = False) -> XlImage | None:
    """whole — диван целиком, с белыми полями; образец ткани обрезается в квадрат."""
    w, h = size
    try:
        with Image.open(path) as src:
            img = ImageOps.exif_transpose(src).convert("RGB")
            img = ImageOps.pad(img, (w * 2, h * 2), color="white") if whole else ImageOps.fit(img, (w * 2, h * 2))
    except OSError:
        return None
    buf = BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    xl = XlImage(buf)
    xl.width, xl.height = w, h
    return xl


def _place(ws, image: XlImage, col: int, row: int, x: int) -> None:
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(x), row=row - 1, rowOff=pixels_to_EMU(4))
    size = XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height))
    image.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(image)


def build_xlsx(result: Result) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", result.title)[:31]
    ws.sheet_view.showGridLines = False
    columns = result.columns
    last = get_column_letter(max(len(columns), 2))

    if LOGO.exists():
        logo = XlImage(str(LOGO))
        logo.width, logo.height = 78, 58
        ws.add_image(logo, "A1")
    ws["B1"] = result.title
    ws["B1"].font = Font(name=FONT, size=15, bold=True, color=INK)
    ws["B2"] = result.subtitle
    ws["B2"].font = Font(name=FONT, size=10, color=MUTED)
    now = timezone.localtime().strftime("%d.%m.%y %H:%M")
    ws["B3"] = f"Сформировано {now}. Всего: {len(result.rows)} шт."
    ws["B3"].font = Font(name=FONT, size=10, color=MUTED)
    ws["B4"] = "Фабрика домашней мебели «Миларум», г. Лесной"
    ws["B4"].font = Font(name=FONT, size=9, color=MUTED)
    for r in range(1, 5):
        ws.merge_cells(f"B{r}:{last}{r}")
    ws.row_dimensions[1].height = 22

    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=col.title)
        cell.font = Font(name=FONT, size=10, bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        width = col.width
        if col.kind == "photos":
            most = max((len(r.photos) for r in result.rows), default=1)
            width = max(8, round((most * (PHOTO + PHOTO_GAP) + 8) / 7))
        if c == 1:
            width = max(width, 19 if any(r.group_photo for r in result.rows) else 12)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[HEADER_ROW].height = 22

    r = HEADER_ROW
    group = None
    for row in result.rows:
        r += 1
        if row.group and row.group != group:
            group = row.group
            image = _photo(row.group_photo.path, GROUP_PHOTO, whole=True) if row.group_photo and row.group_photo.path else None
            first = 2 if image and len(columns) > 1 else 1
            cell = ws.cell(row=r, column=first, value=group)
            cell.font = Font(name=FONT, size=11, bold=True, color=ACCENT)
            cell.alignment = Alignment(vertical="center" if image else "bottom")
            ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=max(len(columns), first))
            ws.row_dimensions[r].height = GROUP_PHOTO[1] * 0.75 + 8 if image else 24
            if image:
                _place(ws, image, 1, r, 4)
            r += 1
        lines = 1
        for c, col in enumerate(columns, start=1):
            value = "" if col.kind == "photos" else row.cells.get(col.key, "")
            if col.kind == "number" and str(value).isdigit():
                value = int(value)
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT, size=10, bold=col.key == "number", color=INK)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row.mark in MARK_FILL:
                cell.fill = PatternFill("solid", fgColor=MARK_FILL[row.mark])
            if isinstance(value, str):
                width = ws.column_dimensions[get_column_letter(c)].width or 10
                lines = max(lines, sum(max(1, -(-len(part) // max(int(width * 0.95), 1))) for part in value.split("\n")))
            if col.kind == "photos":
                x = 4
                for photo in row.photos:
                    image = _photo(photo.path) if photo.path else None
                    if image:
                        _place(ws, image, c, r, x)
                    x += PHOTO + PHOTO_GAP
        has_photos = any(col.kind == "photos" for col in columns) and any(p.path for p in row.photos)
        ws.row_dimensions[r].height = max(lines * 12.5 + 5, PHOTO * 0.75 + 8 if has_photos else 0)

    # фильтр нужен и в отчётах с группами (например, по складу в резервах); заголовки групп при фильтре скрываются.
    # Не ставим, где в заголовках групп картинки: у скрытых строк они наезжали бы на соседние
    if result.rows and not any(row.group_photo for row in result.rows):
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{r}"
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    if result.summary:
        sm = wb.create_sheet("Итог")
        sm.column_dimensions["A"].width = 52
        sm.column_dimensions["B"].width = 10
        line = 1
        for block in result.summary:
            sm.cell(row=line, column=1, value=block.title).font = Font(name=FONT, size=11, bold=True, color=INK)
            line += 1
            for label, value in block.rows:
                sm.cell(row=line, column=1, value=label).font = Font(name=FONT, size=10)
                cell = sm.cell(row=line, column=2, value=int(value) if str(value).isdigit() else value)
                cell.font = Font(name=FONT, size=10)
                line += 1
            line += 1

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
